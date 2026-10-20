"""
ResumeKit backend: static SPA + AI proxy (free OpenRouter models with fallback chain) + DOCX export.
"""
import asyncio
import base64
import io
import json
import os
import re
import time
import uuid
from collections import defaultdict
from datetime import date
from pathlib import Path

import httpx
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from pypdf import PdfReader
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

BASE_DIR = Path(__file__).resolve().parent
STATIC = BASE_DIR / "static"

OPENROUTER_KEY = os.environ.get("OPENROUTER_API_KEY", "")
OPENROUTER_BASE = os.environ.get("OPENROUTER_BASE_URL", "https://openrouter.ai/api/v1")
OPENCODE_KEY = os.environ.get("OPENCODE_GO_API_KEY", "")
OPENCODE_BASE = os.environ.get("OPENCODE_BASE_URL", "https://opencode.ai/zen/go/v1")

# Free-tier OpenRouter models, tried in order until one answers.
DEFAULT_CHAIN = [
    "nvidia/nemotron-3-ultra-550b-a55b:free",
    "minimax/minimax-m3.1-flash:free",
    "nvidia/nemotron-3-super-120b-a12b:free",
    "google/gemma-4-31b-it:free",
    "xingchen-agi/xing4.0-29b-a4b:free",
    "openrouter/free",
]
VISION_MODELS = [
    "nvidia/nemotron-3.5-super-vl:free",
    "openrouter/free",
]
MODELS = [m.strip() for m in os.environ.get("AI_MODELS", "").split(",") if m.strip()] or DEFAULT_CHAIN
MODEL = MODELS[0]
# Parsing a whole resume is a big, mechanical JSON task: the lighter model is
# ~3x faster and just as accurate at extraction, so try it first.
FAST_MODELS = ["minimax/minimax-m3.1-flash:free"] + [m for m in MODELS if "minimax" not in m]

WINDOW = 3600
PER_IP_HOURLY = 60
GLOBAL_DAILY = int(os.environ.get("GLOBAL_DAILY_AI", "1500"))

app = FastAPI(title="ResumeKit", docs_url=None, redoc_url=None)


@app.middleware("http")
async def no_store_html(request: Request, call_next):
    """HTML must never be cached: a stale page ships stale JS to the user."""
    resp = await call_next(request)
    ctype = resp.headers.get("content-type", "")
    if ctype.startswith("text/html") or request.url.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store, max-age=0"
    return resp

_ip_hits = defaultdict(list)
_daily = {"day": date.today().isoformat(), "count": 0}
client = httpx.AsyncClient(timeout=90.0)

# ---------------------------------------------------------------- rate limiting
def check_limit(request: Request):
    global _daily
    today = date.today().isoformat()
    if _daily["day"] != today:
        _daily = {"day": today, "count": 0}
    if _daily["count"] >= GLOBAL_DAILY:
        raise HTTPException(429, "Daily AI quota reached on this demo instance. Try again tomorrow.")
    ip = request.headers.get("cf-connecting-ip") or (request.client.host if request.client else "?")
    now = time.time()
    hits = [t for t in _ip_hits[ip] if now - t < WINDOW]
    if len(hits) >= PER_IP_HOURLY:
        raise HTTPException(429, "You've hit the hourly AI limit. Give it a few minutes.")
    hits.append(now)
    _ip_hits[ip] = hits
    _daily["count"] += 1

# ---------------------------------------------------------------- AI plumbing
def strip_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json|markdown|md|text)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()

def parse_json(text: str):
    text = strip_fences(text)
    try:
        return json.loads(text)
    except Exception:
        pass
    m = re.search(r"(\{.*\}|\[.*\])", text, re.S)
    if m:
        try:
            return json.loads(m.group(1))
        except Exception:
            return None
    return None

def provider_for(model: str):
    """Return (base_url, headers) for a model id."""
    if model.startswith("openrouter/") or model.endswith(":free") or "/" in model:
        if not OPENROUTER_KEY:
            raise HTTPException(500, "OpenRouter key is not configured.")
        return OPENROUTER_BASE, {
            "Authorization": f"Bearer {OPENROUTER_KEY}",
            "Content-Type": "application/json",
            "HTTP-Referer": "https://resume.kavilabs.dev",
            "X-Title": "ResumeKit",
        }
    if not OPENCODE_KEY:
        raise HTTPException(500, "No AI provider is configured for this model.")
    return OPENCODE_BASE, {
        "Authorization": f"Bearer {OPENCODE_KEY}",
        "Content-Type": "application/json",
        "x-opencode-session": str(uuid.uuid4()),
    }


async def _one(model, messages, max_tokens, temperature):
    base, headers = provider_for(model)
    payload = {"model": model, "messages": messages, "max_tokens": max_tokens, "temperature": temperature}
    r = await client.post(f"{base}/chat/completions", json=payload, headers=headers)
    if r.status_code == 429:
        raise HTTPException(429, "rate")
    if r.status_code != 200:
        raise HTTPException(502, f"{model}: HTTP {r.status_code}")
    data = r.json()
    return (data["choices"][0]["message"].get("content") or "").strip()


async def call_ai(messages, max_tokens=900, temperature=0.5, effort="minimal", models=None):
    """Try each free model in the chain. Reasoning models burn budget on thinking,
    so short answers come back empty; retry once with double the budget before
    falling through to the next model."""
    chain = list(models or MODELS)
    last_err = None
    # Free pools get congested in bursts. Sweep the chain three times with
    # backoff before giving up, so a transient 429 doesn't turn into a 502.
    for round_no in range(3):
        for model in chain:
            for attempt, budget in enumerate((max_tokens, max_tokens * 2)):
                try:
                    content = await _one(model, messages, budget, temperature)
                    if content:
                        return content
                except HTTPException as e:
                    last_err = e
                    if e.status_code == 429:
                        await asyncio.sleep(0.8 + round_no * 1.5 + attempt)
                        continue
                    break
                except httpx.TimeoutException:
                    last_err = HTTPException(504, "AI timed out.")
                    break
                except Exception as e:  # malformed provider reply
                    last_err = HTTPException(502, str(e))
                    break
        await asyncio.sleep(1.0 + round_no)
    # Last resort: the flat-rate Go plan, so a user never hits a wall.
    if OPENCODE_KEY and "deepseek-v4.1-flash" not in chain:
        try:
            content = await _one("deepseek-v4.1-flash", messages, max_tokens * 2, temperature)
            if content:
                return content
        except Exception as e:
            last_err = HTTPException(502, str(e))
    raise last_err or HTTPException(502, "All free models are busy right now. Try again in a minute.")


SYSTEM = (
    "You are a senior resume writer and ATS specialist with 15 years of recruiting experience. "
    "You write tight, concrete, metric-driven resume language. No fluff, no first person, no cliches. "
    "Never invent employers or credentials the candidate did not provide. "
    "Follow the requested output format exactly."
)

def resume_digest(resume: dict, limit=3500) -> str:
    parts = []
    c = resume.get("contact", {})
    parts.append(f"Name: {c.get('name','')} | Target title: {c.get('title','')}")
    if resume.get("summary"):
        parts.append(f"Summary: {resume['summary'][:400]}")
    for e in resume.get("experience", [])[:6]:
        parts.append(
            f"EXPERIENCE {e.get('role','')} at {e.get('company','')} ({e.get('start','')}-{e.get('end','')}): "
            + " ".join(e.get("bullets", []))[:600]
        )
    for ed in resume.get("education", [])[:3]:
        parts.append(f"EDUCATION {ed.get('degree','')} {ed.get('field','')} {ed.get('school','')} {ed.get('end','')}")
    skills = resume.get("skills", [])
    flat = []
    for s in skills:
        flat.extend(s.get("items", []) if isinstance(s, dict) else [str(s)])
    if flat:
        parts.append("SKILLS: " + ", ".join(flat)[:400])
    for p in resume.get("projects", [])[:3]:
        parts.append(f"PROJECT {p.get('name','')}: {p.get('description','')}")
    return "\n".join(parts)[:limit]


async def handle_task(task: str, p: dict):
    if task == "bullets":
        role = (p.get("role") or "professional")[:120]
        company = (p.get("company") or "")[:120]
        seed = (p.get("input") or "").strip()[:600]
        n = max(1, min(int(p.get("count") or 4), 5))
        msg = (
            f"Job title: {role}. Company: {company}.\n"
            f"Candidate's rough note: {seed or '(none, infer typical responsibilities)'}\n\n"
            f"Write {n} resume bullet points. Rules: start with a strong past-tense verb, "
            f"include a plausible metric placeholder only if it is clearly derived from the note, "
            f"12-24 words each, no periods at the end, no personal pronouns. "
            f'Return JSON only: {{"bullets": ["...", "..."]}}'
        )
        out = await call_ai([{"role": "system", "content": SYSTEM}, {"role": "user", "content": msg}], 1600)
        data = parse_json(out)
        if isinstance(data, dict) and isinstance(data.get("bullets"), list):
            return {"bullets": [str(b).strip(" -•") for b in data["bullets"]][:n]}
        lines = [l.strip(" -•0123456789.") for l in strip_fences(out).splitlines() if len(l.strip()) > 12]
        return {"bullets": lines[:n]}

    if task == "summary":
        digest = resume_digest(p.get("resume", {}))
        jd = (p.get("jobDescription") or "")[:1500]
        msg = (
            f"Candidate:\n{digest}\n\n" + (f"Target job description:\n{jd}\n\n" if jd else "")
            + "Write a 3-sentence professional summary (45-65 words) in third person without pronouns, "
              "mirroring the target job's keywords where honest. Plain text, no label, no quotes."
        )
        out = await call_ai([{"role": "system", "content": SYSTEM}, {"role": "user", "content": msg}], 1200, 0.55)
        return {"summary": strip_fences(out).strip('"').strip()}

    if task == "skills":
        role = (p.get("role") or "professional")[:120]
        existing = ", ".join(p.get("existing", []) or [])[:400]
        msg = (
            f"Role: {role}. Existing skills: {existing or 'none'}.\n"
            "List 12 hard and soft skills that belong on this resume, ATS-recognizable wording, 1-3 words each. "
            'Return JSON only: {"skills": ["...", "..."]}'
        )
        out = await call_ai([{"role": "system", "content": SYSTEM}, {"role": "user", "content": msg}], 1000)
        data = parse_json(out)
        if isinstance(data, dict) and isinstance(data.get("skills"), list):
            return {"skills": [str(s) for s in data["skills"]][:14]}
        return {"skills": [l.strip(" -•").strip() for l in strip_fences(out).splitlines() if l.strip()][:14]}

    if task == "review":
        digest = resume_digest(p.get("resume", {}))
        msg = (
            f"Review this resume:\n{digest}\n\n"
            "Score it 0-100 for ATS readiness and recruiter impact. Identify up to 5 concrete problems "
            "(weak verbs, missing metrics, length, keyword gaps, formatting risks) and one rewrite for the "
            "single weakest bullet.\n"
            'Return JSON only: {"score": 0, "verdict": "one sentence", "strengths": ["..."], '
            '"fixes": [{"issue": "...", "fix": "..."}], "rewrite": {"before": "...", "after": "..."}}'
        )
        out = await call_ai([{"role": "system", "content": SYSTEM}, {"role": "user", "content": msg}], 2500, 0.35)
        data = parse_json(out)
        if not isinstance(data, dict):
            raise HTTPException(502, "Could not parse the review. Try again.")
        return data

    if task == "cover":
        digest = resume_digest(p.get("resume", {}))
        jd = (p.get("jobDescription") or "")[:2000]
        company = (p.get("company") or "the company")[:100]
        role = (p.get("role") or "the role")[:100]
        msg = (
            f"Candidate:\n{digest}\n\nCompany: {company}\nRole: {role}\nJob description:\n{jd or '(not provided)'}\n\n"
            "Write a cover letter: 3 paragraphs, under 260 words, specific to the role, no cliches "
            "('I am writing to apply'), no invented facts. Start with the hook. Plain text with a greeting "
            "and sign-off using the candidate's name."
        )
        out = await call_ai([{"role": "system", "content": SYSTEM}, {"role": "user", "content": msg}], 2000, 0.6)
        return {"letter": strip_fences(out)}

    if task == "tailor":
        digest = resume_digest(p.get("resume", {}))
        jd = (p.get("jobDescription") or "")[:2000]
        msg = (
            f"Candidate:\n{digest}\n\nJob description:\n{jd}\n\n"
            "Rewrite the summary for this job and propose 3 replacement bullets for the most recent role, "
            "using the JD's exact keywords honestly.\n"
            'Return JSON only: {"summary": "...", "bullets": ["...", "...", "..."], '
            '"keywords": ["..."]}'
        )
        out = await call_ai([{"role": "system", "content": SYSTEM}, {"role": "user", "content": msg}], 2200, 0.45)
        data = parse_json(out)
        if not isinstance(data, dict):
            raise HTTPException(502, "Could not parse tailoring output. Try again.")
        return data

    raise HTTPException(400, "Unknown task")

@app.post("/api/ai")
async def ai_endpoint(request: Request):
    check_limit(request)
    body = await request.json()
    task = body.get("task")
    payload = body.get("payload") or {}
    return JSONResponse(await handle_task(task, payload))

@app.get("/api/health")
async def health():
    return {"ok": True, "models": MODELS, "providers": {"openrouter": bool(OPENROUTER_KEY), "opencode": bool(OPENCODE_KEY)}}

# ---------------------------------------------------------------- DOCX export
def _p(doc, text="", size=10.5, bold=False, italic=False, space_after=2, space_before=0,
       align=None, color=None, upper=False):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.06
    if align:
        par.alignment = align
    run = par.add_run(text.upper() if upper else text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = RGBColor(*color)
    return par

def build_docx(r: dict) -> io.BytesIO:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.55)
        s.bottom_margin = Inches(0.55)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    c = r.get("contact", {})
    if c.get("name"):
        _p(doc, c["name"], 20, True, space_after=1)
    if c.get("title"):
        _p(doc, c["title"], 11.5, False, italic=True, space_after=3, color=(0x33, 0x33, 0x33))
    contact = " | ".join(x for x in [c.get("email"), c.get("phone"), c.get("location"),
                                     c.get("website"), c.get("linkedin")] if x)
    if contact:
        _p(doc, contact, 9.5, space_after=6, color=(0x44, 0x44, 0x44))

    def heading(t):
        par = _p(doc, t, 11, True, space_before=8, space_after=3, upper=True, color=(0x11, 0x11, 0x11))
        pr = par._p.get_or_add_pPr()
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "999999")
        bdr.append(bottom)
        pr.append(bdr)

    if r.get("summary"):
        heading("Summary")
        _p(doc, r["summary"], 10.5, space_after=2)

    if r.get("experience"):
        heading("Experience")
        for e in r["experience"]:
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.space_before = Pt(4); pf.space_after = Pt(0)
            pf.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
            run = par.add_run(e.get("role", ""))
            run.bold = True; run.font.size = Pt(10.5); run.font.name = "Calibri"
            dates = " – ".join(x for x in [e.get("start"), "Present" if e.get("current") else e.get("end")] if x)
            if dates:
                par.add_run("\t" + dates).font.size = Pt(9.5)
            sub = " — ".join(x for x in [e.get("company"), e.get("location")] if x)
            if sub:
                _p(doc, sub, 10, italic=True, space_after=1, color=(0x33, 0x33, 0x33))
            for b in e.get("bullets", []):
                if b.strip():
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.space_after = Pt(1)
                    bp.paragraph_format.line_spacing = 1.04
                    bp.add_run(b.strip()).font.size = Pt(10.5)

    if r.get("education"):
        heading("Education")
        for ed in r["education"]:
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(3); par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
            line = ", ".join(x for x in [ed.get("degree"), ed.get("field")] if x) or "Degree"
            run = par.add_run(line); run.bold = True; run.font.size = Pt(10.5)
            dates = " – ".join(x for x in [ed.get("start"), ed.get("end")] if x)
            if dates:
                par.add_run("\t" + dates).font.size = Pt(9.5)
            sub = " — ".join(x for x in [ed.get("school"), ed.get("location")] if x)
            if sub:
                _p(doc, sub, 10, italic=True, space_after=1, color=(0x33, 0x33, 0x33))
            if ed.get("details"):
                _p(doc, ed["details"], 10, space_after=1)

    if r.get("projects"):
        heading("Projects")
        for pr in r["projects"]:
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(3); par.paragraph_format.space_after = Pt(0)
            run = par.add_run(pr.get("name", "")); run.bold = True; run.font.size = Pt(10.5)
            if pr.get("link"):
                par.add_run("  " + pr["link"]).font.size = Pt(9.5)
            if pr.get("description"):
                _p(doc, pr["description"], 10, space_after=1)

    skills = r.get("skills", [])
    if skills:
        heading("Skills")
        for s in skills:
            items = s.get("items", []) if isinstance(s, dict) else [s]
            if not items:
                continue
            label = (s.get("category") or "").strip() if isinstance(s, dict) else ""
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(1)
            if label:
                run = par.add_run(label + ": "); run.bold = True; run.font.size = Pt(10.5)
            par.add_run(", ".join(items)).font.size = Pt(10.5)

    if r.get("certifications"):
        heading("Certifications")
        for cert in r["certifications"]:
            line = " — ".join(x for x in [cert.get("name"), cert.get("issuer"), cert.get("year")] if x)
            _p(doc, line, 10, space_after=1)

    if r.get("languages"):
        heading("Languages")
        _p(doc, ", ".join(f"{l.get('name','')} ({l.get('level','')})" if l.get("level") else l.get("name", "")
                          for l in r["languages"]), 10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

@app.post("/api/export/docx")
async def export_docx(request: Request):
    body = await request.json()
    resume = body.get("resume") or {}
    buf = build_docx(resume)
    name = re.sub(r"[^A-Za-z0-9]+", "_", (resume.get("contact", {}).get("name") or "resume")).strip("_") or "resume"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}_Resume.docx"'},
    )

# ---------------------------------------------------------------- resume import
SCHEMA_HINT = """{
  "contact": {"name": "", "title": "", "email": "", "phone": "", "location": "", "website": "", "linkedin": ""},
  "summary": "",
  "experience": [{"role": "", "company": "", "location": "", "start": "", "end": "", "current": false, "bullets": [""]}],
  "education": [{"degree": "", "field": "", "school": "", "location": "", "start": "", "end": "", "gpa": "", "details": ""}],
  "skills": [{"category": "", "items": [""]}],
  "projects": [{"name": "", "link": "", "description": ""}],
  "certifications": [{"name": "", "issuer": "", "year": ""}],
  "languages": [{"name": "", "level": ""}]
}"""


def clean_text(raw: str) -> str:
    raw = raw.replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in raw.split("\n")]
    out, blank = [], False
    for ln in lines:
        if not ln:
            if not blank:
                out.append("")
            blank = True
        else:
            out.append(ln)
            blank = False
    return "\n".join(out).strip()


def pdf_text(blob: bytes) -> str:
    reader = PdfReader(io.BytesIO(blob))
    chunks = []
    for page in reader.pages[:12]:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    return clean_text("\n".join(chunks))


def pdf_images(blob: bytes, max_pages: int = 3) -> list:
    """Render pages to base64 PNGs for scanned/image-only PDFs."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return []
    out = []
    doc = fitz.open(stream=blob, filetype="pdf")
    for page in list(doc)[:max_pages]:
        pix = page.get_pixmap(dpi=150)
        out.append(base64.b64encode(pix.tobytes("png")).decode())
    return out


def docx_text(blob: bytes) -> str:
    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return clean_text("\n".join(parts))


def s(v, limit=400) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()[:limit]


def as_list(v):
    if isinstance(v, list):
        return v
    if v in (None, ""):
        return []
    return [v]


def normalize(parsed: dict) -> dict:
    """Coerce whatever the model returned into the builder's shape."""
    if not isinstance(parsed, dict):
        raise HTTPException(502, "Could not read that resume. Try the paste option below.")
    if isinstance(parsed.get("resume"), dict):
        parsed = parsed["resume"]

    c = parsed.get("contact") or {}
    contact = {k: s(c.get(k), 200) for k in
               ("name", "title", "email", "phone", "location", "website", "linkedin")}
    # Resumes are often set in all caps; the builder wants normal casing.
    if contact["name"] and contact["name"].isupper():
        contact["name"] = contact["name"].title()

    experience = []
    for e in as_list(parsed.get("experience"))[:12]:
        if not isinstance(e, dict):
            continue
        bullets = [s(b, 400) for b in as_list(e.get("bullets") or e.get("achievements")) if s(b)]
        if not bullets and e.get("description"):
            bullets = [s(x, 400) for x in re.split(r"(?<=[.!?])\s+|\n", s(e["description"], 2000)) if len(x) > 12]
        experience.append({
            "role": s(e.get("role") or e.get("title") or e.get("position"), 160),
            "company": s(e.get("company") or e.get("employer") or e.get("organization"), 160),
            "location": s(e.get("location"), 120),
            "start": s(e.get("start") or e.get("start_date"), 60),
            "end": s(e.get("end") or e.get("end_date"), 60),
            "current": bool(e.get("current")) or s(e.get("end")).lower() in ("present", "current", "now"),
            "bullets": bullets,
        })

    education = []
    for e in as_list(parsed.get("education"))[:8]:
        if not isinstance(e, dict):
            continue
        education.append({
            "degree": s(e.get("degree"), 120),
            "field": s(e.get("field") or e.get("major"), 120),
            "school": s(e.get("school") or e.get("institution"), 160),
            "location": s(e.get("location"), 120),
            "start": s(e.get("start") or e.get("start_date"), 60),
            "end": s(e.get("end") or e.get("end_date"), 60),
            "gpa": s(e.get("gpa"), 40),
            "details": s(e.get("details") or e.get("honors"), 300),
        })

    skills = []
    for grp in as_list(parsed.get("skills"))[:8]:
        if isinstance(grp, dict):
            items = [s(i, 60) for i in as_list(grp.get("items") or grp.get("skills")) if s(i)]
            if items:
                skills.append({"category": s(grp.get("category") or grp.get("name") or "Skills", 60), "items": items})
        elif isinstance(grp, str) and s(grp):
            skills.append({"category": "Skills", "items": [x.strip() for x in re.split(r"[,;|/]", grp) if x.strip()][:30]})
    if not skills:
        plain = parsed.get("skills")
        if isinstance(plain, str) and plain.strip():
            skills.append({"category": "Skills", "items": [x.strip() for x in re.split(r"[,;|]", plain) if x.strip()]})

    projects = []
    for pr in as_list(parsed.get("projects"))[:8]:
        if isinstance(pr, dict):
            projects.append({"name": s(pr.get("name") or pr.get("title"), 160),
                             "link": s(pr.get("link") or pr.get("url"), 200),
                             "description": s(pr.get("description") or pr.get("details"), 600)})

    certs = []
    for c2 in as_list(parsed.get("certifications"))[:10]:
        if isinstance(c2, dict):
            certs.append({"name": s(c2.get("name") or c2.get("certification"), 160),
                          "issuer": s(c2.get("issuer") or c2.get("organization"), 160),
                          "year": s(c2.get("year") or c2.get("date"), 40)})
        elif isinstance(c2, str):
            certs.append({"name": s(c2, 160), "issuer": "", "year": ""})

    langs = []
    for l in as_list(parsed.get("languages"))[:8]:
        if isinstance(l, dict):
            langs.append({"name": s(l.get("name") or l.get("language"), 80), "level": s(l.get("level") or l.get("proficiency"), 60)})
        elif isinstance(l, str):
            langs.append({"name": s(l, 80), "level": ""})

    out = {
        "contact": contact,
        "summary": s(parsed.get("summary") or parsed.get("profile") or parsed.get("objective"), 1200),
        "experience": [e for e in experience if e["role"] or e["company"]],
        "education": [e for e in education if e["school"] or e["degree"]],
        "skills": skills,
        "projects": [p for p in projects if p["name"] or p["description"]],
        "certifications": [c3 for c3 in certs if c3["name"]],
        "languages": [l for l in langs if l["name"]],
    }
    if not any([out["contact"]["name"], out["experience"], out["education"], out["summary"]]):
        raise HTTPException(422, "I couldn't find resume content in that file. Try pasting the text instead.")
    return out


async def parse_resume_text(text: str) -> dict:
    text = clean_text(text)[:16000]
    if len(text) < 60:
        raise HTTPException(422, "That file has almost no text in it. It may be a scan or an image-only PDF.")
    messages = [
        {"role": "system", "content": SYSTEM + " You extract resume data into JSON exactly, without inventing anything."},
        {"role": "user", "content": (
            "Extract this resume into JSON. Copy the wording as written; do not rewrite, improve, or invent. "
            "Split the experience section into one object per job with its bullets. "
            "Keep dates in the form they appear. Use \"current\": true when the job says Present. "
            "Group skills under sensible categories. Omit anything absent.\n\n"
            f"JSON shape:\n{SCHEMA_HINT}\n\nRESUME TEXT:\n{text}"
        )},
    ]
    raw = await call_ai(messages, 3000, 0.15, models=FAST_MODELS)
    data = parse_json(raw)
    if data is None:
        raise HTTPException(502, "The model returned unreadable data. Try again or paste the text.")
    return normalize(data)


async def parse_resume_images(images: list) -> dict:
    content = [{"type": "text", "text": (
        "This is a scan of a resume. Read every word and extract it into JSON exactly as written; "
        "do not rewrite or invent. One object per job with its bullets. current=true when the job says Present. "
        "Group skills under sensible categories. Omit anything absent.\n\n"
        f"JSON shape:\n{SCHEMA_HINT}"
    )}]
    for b64 in images:
        content.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}})
    raw = await call_ai([{"role": "user", "content": content}], 3000, 0.15,
                        models=VISION_MODELS + [m for m in FAST_MODELS if m not in VISION_MODELS])
    data = parse_json(raw)
    if data is None:
        raise HTTPException(502, "Could not read that scan. Try pasting the text instead.")
    return normalize(data)


@app.post("/api/import")
async def import_resume(request: Request, file: UploadFile = File(...)):
    check_limit(request)
    blob = await file.read()
    if len(blob) > 8 * 1024 * 1024:
        raise HTTPException(413, "That file is larger than 8 MB.")
    name = (file.filename or "").lower()
    try:
        if name.endswith(".pdf") or blob[:4] == b"%PDF":
            text = pdf_text(blob)
        elif name.endswith((".docx", ".doc")):
            try:
                text = docx_text(blob)
            except Exception:
                raise HTTPException(422, "Could not read that Word file. Save it as PDF and try again.")
        elif name.endswith((".txt", ".md")):
            text = clean_text(blob.decode("utf-8", "ignore"))
        else:
            raise HTTPException(415, "Upload a PDF, Word file, or plain text resume.")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, f"Could not read that file: {e}")

    # Image-only PDF: no text layer, so read the pages visually instead.
    if len(text) < 150 and (name.endswith(".pdf") or blob[:4] == b"%PDF"):
        pages = pdf_images(blob)
        if pages:
            return JSONResponse({"resume": await parse_resume_images(pages), "method": "vision"})
        raise HTTPException(422, "That PDF is a scan with no text layer. Paste the text instead.")

    return JSONResponse({"resume": await parse_resume_text(text)})


@app.post("/api/import-text")
async def import_resume_text(request: Request):
    check_limit(request)
    body = await request.json()
    return JSONResponse({"resume": await parse_resume_text(body.get("text") or "")})


# ---------------------------------------------------------------- static
@app.get("/build")
async def build_page():
    return FileResponse(STATIC / "build.html")

@app.get("/")
async def index():
    return FileResponse(STATIC / "index.html")

app.mount("/", StaticFiles(directory=str(STATIC), html=True), name="static")
